# -*- coding: utf-8 -*-
"""把 OHR-Bench 物化成演示知识库 `rag_data/`：**多文档、多格式**。

为什么必须多格式：校内初赛「技术数据」硬性要求
「知识库文档样本（不少于 10 份，涵盖不同类型文档格式）」。
同时这也让 RAG 演示更像真实企业场景（赛题原话：PDF 报告、技术文档、规章制度）。

数据来源：OHR-Bench 的 `gt` 检索基（官方推荐数据集，7 领域 1261 份文档）。
输出格式：PDF / DOCX / PPTX / XLSX / CSV / MD / TXT —— 覆盖文本、表格、演示、报告四类形态。

注意：`RAGTool._load_documents()` **只扫 rag_data 顶层**（不递归），所以全部平铺命名。
表格类格式（XLSX/CSV）优先挑原文含 LaTeX tabular 的页，解析成真正的行列再落盘，
而不是把整段文字塞进一个单元格。

用法：
    python scripts/build_demo_kb.py            # 生成
    python scripts/build_demo_kb.py --clean    # 先清掉上一次生成的
"""
import argparse
import json
import os
import re
import sys
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

WS = r"D:\硬件数据\deepseek-harness-project\中电子云"
OHR = os.path.join(WS, "data", "ohr")
ZIP = os.path.join(OHR, "retrieval.zip")
KB = os.path.join(ROOT, "rag_data")
# 清单**不能放在 rag_data 里**：即使有 db_builder 的下划线过滤兜底，
# 知识库目录也只应该装知识文档本身。放到 data/ 下更干净。
MANIFEST = os.path.join(ROOT, "data", "kb_manifest.json")

PAGES_PER_DOC = 8

# (输出格式, 领域, 期望页内含表格) —— 顺序即执行顺序，保证可复现
PLAN = [
    ("pdf",  "finance",  False),
    ("pdf",  "law",      False),
    ("docx", "law",      False),
    ("docx", "textbook", False),
    ("pptx", "manual",   False),
    ("xlsx", "finance",  True),
    ("xlsx", "law",      True),
    ("csv",  "finance",  True),
    ("md",   "textbook", False),
    ("md",   "paper",    False),
    ("txt",  "news",     False),
    ("pdf",  "manual",   False),
]

TABULAR = re.compile(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}", re.S)
_ROW_NOISE = re.compile(r"\\(hline|toprule|midrule|bottomrule|endhead|endfirsthead)\b")
_CLINE = re.compile(r"\\cline\{[^}]*\}")
_SPAN = re.compile(r"\\(?:multi|)column\{[^}]*\}\{[^}]*\}|\\multirow\{[^}]*\}\{[^}]*\}")


def parse_tabular(text: str):
    """把一个 LaTeX tabular 解析成 rows（list[list[str]]），找不到返回 []"""
    m = TABULAR.search(text or "")
    if not m:
        return []
    rows = []
    for raw in re.split(r"\\\\", m.group(1)):
        line = _CLINE.sub("", _ROW_NOISE.sub("", raw)).strip()
        if not line:
            continue
        cells = [_SPAN.sub("", c).strip() for c in line.split("&")]
        cells = [c[1:-1].strip() if c.startswith("{") and c.endswith("}") else c
                 for c in cells]
        if any(cells):
            rows.append(cells)
    if len(rows) < 2:
        return []
    width = max(len(r) for r in rows)
    return [r + [""] * (width - len(r)) for r in rows]


def safe(name: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff.\-]+", "_", name)[:80]


def load_docs():
    """从 zip 里按领域读 gt 页文本，返回 {doc_name: [(page_idx, text)]}"""
    out = {}
    with zipfile.ZipFile(ZIP) as z:
        for n in sorted(z.namelist()):
            if not (n.startswith("gt/") and n.endswith(".json")):
                continue
            parts = n.split("/")
            if len(parts) != 3:
                continue
            doc = f"{parts[1]}/{parts[2][:-5]}"
            with z.open(n) as fh:
                pages = json.loads(fh.read().decode("utf-8", "replace"))
            pages = [(int(p.get("page_idx", 0)), (p.get("text") or "").strip())
                     for p in (pages or []) if (p.get("text") or "").strip()]
            if pages:
                out[doc] = sorted(pages)
    return out


def pick(docs, domain, want_table, used):
    """在同一领域里挑一份还没用过的文档（要表格就优先含 tabular 的）"""
    cands = [d for d in sorted(docs) if d.startswith(domain + "/") and d not in used]
    if want_table:
        cands.sort(key=lambda d: -sum(1 for _, t in docs[d] if TABULAR.search(t)))
    return cands[0] if cands else None


# ---------------- 各格式写出 ----------------
def w_txt(path, title, pages):
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"{title}\n{'=' * 60}\n\n")
        for i, t in pages:
            f.write(f"[page {i}]\n{t}\n\n")


def w_md(path, title, pages):
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# {title}\n\n")
        for i, t in pages:
            f.write(f"## Page {i}\n\n{t}\n\n")


def w_pdf(path, title, pages):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=8.5, leading=11.5,
                          wordWrap="CJK")
    head = ParagraphStyle("head", parent=ss["Heading1"], fontSize=14)
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm, title=title)
    flow = [Paragraph(title, head), Spacer(1, 6)]
    for i, t in pages:
        flow.append(Paragraph(f"Page {i}", ss["Heading3"]))
        for para in t.split("\n\n"):
            # 顺序很重要：**先截断、再做实体转义、最后把换行换成 <br/>**。
            # 反过来的话（先插 <br/> 再截断）会把标签从中间切断，
            # reportlab 的 paraparser 就会报 "unclosed tags br/para"（踩过）。
            raw = para[:3000]
            esc = (raw.replace("&", "&amp;").replace("<", "&lt;")
                   .replace(">", "&gt;").replace("\n", "<br/>"))
            if esc.strip():
                flow.append(Paragraph(esc, body))
        flow.append(PageBreak())
    doc.build(flow)


def w_docx(path, title, pages):
    from docx import Document
    d = Document()
    d.add_heading(title, level=0)
    for i, t in pages:
        d.add_heading(f"Page {i}", level=2)
        for para in t.split("\n\n"):
            if para.strip():
                d.add_paragraph(para[:5000])
    d.save(path)


def w_pptx(path, title, pages):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = f"{len(pages)} pages · OHR-Bench 公开数据集"
    for i, t in pages:
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = f"Page {i}"
        tf = sl.placeholders[1].text_frame
        tf.text = t[:600]
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(11)
    prs.save(path)


def w_xlsx(path, title, pages):
    import pandas as pd
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        wrote = 0
        for i, t in pages:
            rows = parse_tabular(t)
            if len(rows) < 2:
                continue
            hdr = [str(c) or f"列{j+1}" for j, c in enumerate(rows[0])]
            df = pd.DataFrame(rows[1:], columns=hdr)
            sheet = f"page{i}"[:31]
            df.to_excel(xw, index=False, sheet_name=sheet)
            wrote += 1
        if not wrote:
            pd.DataFrame({"说明": [title], "内容": ["该文档无可解析表格"]}).to_excel(
                xw, index=False, sheet_name="Sheet0")


def w_csv(path, title, pages):
    import csv
    for i, t in pages:
        rows = parse_tabular(t)
        if len(rows) >= 2:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerows(rows)
            return
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows([["说明", "该文档无可解析表格"], [title, ""]])


WRITERS = {"txt": w_txt, "md": w_md, "pdf": w_pdf,
           "docx": w_docx, "pptx": w_pptx, "xlsx": w_xlsx, "csv": w_csv}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--clean", action="store_true")
    args = ap.parse_args()

    if not os.path.exists(ZIP):
        print(f"❌ 缺少 {ZIP}")
        return 1
    os.makedirs(KB, exist_ok=True)
    if args.clean and os.path.exists(MANIFEST):
        for f in json.load(open(MANIFEST, encoding="utf-8")).get("files", []):
            p = os.path.join(KB, f)
            if os.path.exists(p):
                os.remove(p)
        print("已清理上一次生成的文件")

    docs = load_docs()
    print(f"可用文档 {len(docs)} 份")
    used, made = set(), []
    for fmt, domain, want_table in PLAN:
        doc = pick(docs, domain, want_table, used)
        if not doc:
            print(f"  ⚠️ {domain} 没有可用文档，跳过 {fmt}")
            continue
        used.add(doc)
        pages = docs[doc][:PAGES_PER_DOC]
        # 表格类格式：把含 tabular 的页排到前面，保证能解析出行列
        if want_table:
            pages = sorted(pages, key=lambda pt: 0 if TABULAR.search(pt[1]) else 1)
        title = doc.split("/", 1)[1].replace("_", " ")
        fn = f"{safe(doc.replace('/', '_'))}.{fmt}"
        path = os.path.join(KB, fn)
        try:
            WRITERS[fmt](path, title, pages)
            size = os.path.getsize(path)
            made.append(fn)
            print(f"  ✅ {fmt:5s} {fn[:68]:70s} {size/1024:8.1f} KB  ({domain})")
        except Exception as e:
            print(f"  ❌ {fmt:5s} {fn}: {type(e).__name__}: {str(e)[:120]}")

    # 已有文件（如用户的综测 xlsx / PDF）不动
    existing = [f for f in sorted(os.listdir(KB))
                if os.path.isfile(os.path.join(KB, f)) and not f.startswith("_")]
    with open(MANIFEST, "w", encoding="utf-8") as f:
        json.dump({"generated_by": "scripts/build_demo_kb.py",
                   "source": "OHR-Bench gt retrieval base",
                   "files": made}, f, ensure_ascii=False, indent=2)

    exts = {}
    for f in existing:
        exts[os.path.splitext(f)[1].lower()] = exts.get(os.path.splitext(f)[1].lower(), 0) + 1
    print()
    print(f"✅ rag_data 现有 {len(existing)} 份文件，格式分布：{exts}")
    print("   硬性要求：≥10 份且涵盖不同格式 →",
          "满足" if len(existing) >= 10 and len(exts) >= 3 else "不满足")
    print("   ⚠️ 新增文件会改变索引指纹，下次启动会重建向量索引（需要 embedding 接口可用）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
